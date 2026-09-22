"""
ui/cost_change_log_dialog.py
"Cost change log": every recorded unit-cost change, newest first, so price
movements can be spotted instead of being buried in the item table.

Data comes straight from the audit log, which already records one entry per
changed field (old value / new value / timestamp).
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from PySide6.QtWidgets import (
    QAbstractItemView, QComboBox, QDialog, QFileDialog, QHBoxLayout, QLabel,
    QMessageBox, QPushButton, QTableWidgetItem, QVBoxLayout,
)
from PySide6.QtCore import Qt

from i18n import tr
from ui.table_utils import EmptyStateTable, configure_interactive_table

COLUMNS = ["Date", "Project", "Item", "Old cost", "New cost", "Change %"]

#: groups the report dialog offers, matching get_cost_changes' scope.
SCOPES = [("project", "var_scope_project"), ("all", "var_scope_all")]


def _number(text):
    """old_value/new_value are stored as text; returns float or None."""
    try:
        return float(str(text))
    except (TypeError, ValueError):
        return None


def change_percent(old_text, new_text):
    """Signed percentage of the move, or None when it cannot be computed."""
    old, new = _number(old_text), _number(new_text)
    if old in (None, 0) or new is None:
        return None
    return (new - old) / old * 100.0


def _change_label(old_text, new_text):
    percent = change_percent(old_text, new_text)
    if percent is None:
        return "-"
    arrow = "▲" if percent > 0 else ("▼" if percent < 0 else "=")
    return f"{arrow} {abs(percent):.0f}%"


class CostChangeLogDialog(QDialog):
    def __init__(self, parent, db, project_id=None, project_name=None):
        super().__init__(parent)
        self.db = db
        self.project_id = project_id
        self.project_name = project_name
        self.rows = []

        self.setWindowTitle(tr("cost_title"))
        self.setMinimumSize(920, 560)

        layout = QVBoxLayout(self)
        header = QLabel(
            f"{tr('var_scope_label')}: {project_name}" if project_name
            else f"{tr('var_scope_label')}: {tr('var_scope_all')}"
        )
        header.setObjectName("breadcrumb")
        layout.addWidget(header)

        controls = QHBoxLayout()
        controls.addWidget(QLabel(tr("var_scope_label")))
        self.scope_combo = QComboBox()
        self.scope_combo.addItem(tr("var_scope_project"), "project")
        self.scope_combo.addItem(tr("var_scope_all"), "all")
        self.scope_combo.currentIndexChanged.connect(self._reload)
        controls.addWidget(self.scope_combo)
        controls.addStretch()
        layout.addLayout(controls)

        self.table = EmptyStateTable(tr("cost_empty"))
        self.table.setColumnCount(len(COLUMNS))
        self.table.setHorizontalHeaderLabels(COLUMNS)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setAlternatingRowColors(True)
        configure_interactive_table(self.table, first_col_width=190, layout_key="cost_changes")
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeaderItem(0).setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.table.horizontalHeaderItem(3).setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.table.horizontalHeaderItem(4).setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.table.horizontalHeaderItem(5).setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
        layout.addWidget(self.table)

        buttons = QHBoxLayout()
        export_btn = QPushButton(tr("export_excel"))
        export_btn.setObjectName("primaryButton")
        export_btn.clicked.connect(self._export_excel)
        buttons.addWidget(export_btn)
        buttons.addStretch()
        close_btn = QPushButton(tr("close"))
        close_btn.clicked.connect(self.accept)
        buttons.addWidget(close_btn)
        layout.addLayout(buttons)

        self._reload()

    # ------------------------------------------------------------------ #
    def _reload(self):
        scope = self.scope_combo.currentData() if self.scope_combo else "all"
        self.rows = self.db.get_cost_changes(
            None if scope == "all" or self.project_id is None else [self.project_id]
        )

        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(self.rows))
        for row, entry in enumerate(self.rows):
            values = [
                entry["changed_at"] or "",
                entry["project_name"] or "-",
                entry["item_name"] or "-",
                entry["old_value"] if entry["old_value"] is not None else "-",
                entry["new_value"] if entry["new_value"] is not None else "-",
                _change_label(entry["old_value"], entry["new_value"]),
            ]
            for col, value in enumerate(values):
                cell = QTableWidgetItem(str(value))
                if col >= 3:
                    cell.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(row, col, cell)
        self.table.setSortingEnabled(True)
        self.table.resizeRowsToContents()

    # ------------------------------------------------------------------ #
    def _export_excel(self):
        if not self.rows:
            QMessageBox.information(self, tr("cost_title"), tr("var_empty"))
            return
        scope = (f"{tr('var_scope_label')}: "
                 + (self.project_name if self.project_id is not None else tr("var_scope_all")))
        path, _ = QFileDialog.getSaveFileName(self, tr("export_excel"),
                                              "cost_changes.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        try:
            write_cost_changes_excel(self.rows, path, scope)
        except Exception as exc:  # noqa: BLE001
            QMessageBox.critical(self, tr("export_excel"), str(exc))
            return
        QMessageBox.information(self, tr("export_excel"), f"{tr('export_done')}\n{path}")


def write_cost_changes_excel(rows, output_path, scope_label):
    """One styled sheet: date, project, item, old cost, new cost, change %."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cost changes"

    last_col = get_column_letter(len(COLUMNS))
    sheet.merge_cells(f"A1:{last_col}1")
    title = sheet.cell(row=1, column=1, value=f"{tr('cost_title')} - {scope_label}")
    title.font = Font(name="Calibri", bold=True, size=14)

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    header_font = Font(name="Calibri", bold=True, size=10)
    for index, label in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=4, column=index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = sheet.cell(row=5, column=1)

    for offset, entry in enumerate(rows):
        row_index = 5 + offset
        old_text = entry["old_value"]
        new_text = entry["new_value"]
        values = [
            entry["changed_at"] or "",
            entry["project_name"] or "-",
            entry["item_name"] or "-",
            _number(old_text) if _number(old_text) is not None else (old_text or "-"),
            _number(new_text) if _number(new_text) is not None else (new_text or "-"),
            _change_label(old_text, new_text),
        ]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=index, value=value)
            cell.font = Font(name="Calibri", size=10)
            if index >= 4:
                cell.alignment = Alignment(horizontal="right")
                if index in (4, 5) and isinstance(value, float):
                    cell.number_format = "#,##0.00"

    widths = (20, 24, 40, 14, 14, 12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(output_path)
    return output_path
