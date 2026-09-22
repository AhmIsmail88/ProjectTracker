"""
ui/variance_report_dialog.py
"Variance report": supply status aggregated by area, supplier or status, so a
QS can see where a project stands rather than reading item by item.

Quantities are shown as counts and percentages, never summed: items inside one
group can use different units (عدد / م.ط / مجموعة), so a total would be
meaningless. Value is grouped per currency for the same reason.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from PySide6.QtWidgets import (
    QAbstractItemView, QComboBox, QDialog, QFileDialog, QHBoxLayout, QHeaderView,
    QLabel, QMessageBox, QPushButton, QTableWidgetItem, QVBoxLayout,
)
from PySide6.QtCore import Qt

from i18n import tr
from ui.table_utils import EmptyStateTable, configure_interactive_table

COLUMNS = ["Group", "Items", "Delivered", "Delivered %", "Awaiting request",
           "Needs attention", "Value"]

GROUPINGS = [("area", "var_group_area"), ("supplier", "var_group_supplier"),
             ("status", "var_group_status")]


def _money(values, primary, mixed):
    """'-' when nothing is priced, a single figure when one currency is in
    play, and the primary figure plus a marker when currencies are mixed."""
    if not values or not primary:
        return "-"
    text = f"{values[primary]:,.2f} {primary}"
    return f"{text} *" if mixed else text


class VarianceReportDialog(QDialog):
    def __init__(self, parent, db, project_id=None, project_name=None):
        super().__init__(parent)
        self.db = db
        self.project_id = project_id
        self.project_name = project_name

        self.setWindowTitle(tr("var_title"))
        self.setMinimumSize(920, 560)

        layout = QVBoxLayout(self)

        header = QLabel(
            f"{tr('var_scope_label')}: {project_name}" if project_name
            else f"{tr('var_scope_label')}: {tr('var_scope_all')}"
        )
        header.setObjectName("breadcrumb")
        layout.addWidget(header)

        controls = QHBoxLayout()
        controls.addWidget(QLabel(tr("var_group_by")))
        self.group_combo = QComboBox()
        for key, label_key in GROUPINGS:
            self.group_combo.addItem(tr(label_key), key)
        self.group_combo.currentIndexChanged.connect(self._reload)
        controls.addWidget(self.group_combo)

        if project_id is not None:
            controls.addSpacing(16)
            controls.addWidget(QLabel(tr("var_scope_label")))
            self.scope_combo = QComboBox()
            self.scope_combo.addItem(tr("var_scope_project"), "project")
            self.scope_combo.addItem(tr("var_scope_all"), "all")
            self.scope_combo.currentIndexChanged.connect(self._reload)
            controls.addWidget(self.scope_combo)
        else:
            self.scope_combo = None

        controls.addStretch()
        layout.addLayout(controls)

        self.table = EmptyStateTable(tr("var_empty"))
        self.table.setColumnCount(len(COLUMNS))
        self.table.setHorizontalHeaderLabels(COLUMNS)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setAlternatingRowColors(True)
        configure_interactive_table(self.table, first_col_width=200, layout_key="variance_report")
        self.table.verticalHeader().setVisible(False)
        layout.addWidget(self.table)

        note = QLabel(tr("var_note"))
        note.setObjectName("breadcrumb")
        note.setWordWrap(True)
        layout.addWidget(note)

        buttons = QHBoxLayout()
        export_btn = QPushButton(tr("export_excel"))
        export_btn.setObjectName("primaryButton")
        export_btn.clicked.connect(self._export_excel)
        buttons.addWidget(export_btn)
        buttons.addStretch()
        close_btn = QPushButton(tr("close"))
        close_btn.clicked.connect(self.accept)
        buttons.addWidget(close_btn)
        layout.addLayout(buttons)

        self._reload()

    # ------------------------------------------------------------------ #
    def _current_scope_ids(self):
        if self.scope_combo is None or self.scope_combo.currentData() == "project":
            return None if self.scope_combo is None else {self.project_id}
        return None

    def _reload(self):
        grouping = self.group_combo.currentData()
        self.rows = self.db.get_variance_report(grouping, self._current_scope_ids())

        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(self.rows))
        for row, data in enumerate(self.rows):
            values = [
                data["group"],
                str(data["items"]),
                str(data["delivered"]),
                f"{data['delivered_pct']:.0f}%",
                str(data["awaiting_request"]),
                str(data["attention"]),
                _money(data["values"], data["currency"], data["mixed_currency"]),
            ]
            for col, value in enumerate(values):
                cell = QTableWidgetItem(value)
                if col >= 1:
                    cell.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.table.setItem(row, col, cell)
        self.table.setSortingEnabled(True)
        self.table.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.Interactive)
        self.table.resizeRowsToContents()

    # ------------------------------------------------------------------ #
    def _export_excel(self):
        if not self.rows:
            QMessageBox.information(self, tr("var_title"), tr("var_empty"))
            return

        scope = (f"{tr('var_scope_label')}: "
                 + (self.project_name if self.current_scope_is_project() else tr("var_scope_all")))
        suggested = f"variance_{self.group_combo.currentData()}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, tr("export_excel"), suggested, "Excel (*.xlsx)")
        if not path:
            return

        try:
            write_variance_excel(self.rows, path, self.group_combo.currentText(), scope)
        except Exception as exc:  # noqa: BLE001 - report it, never crash
            QMessageBox.critical(self, tr("export_excel"), str(exc))
            return

        QMessageBox.information(self, tr("export_excel"), f"{tr('export_done')}\n{path}")

    def current_scope_is_project(self):
        return self.scope_combo is None or self.scope_combo.currentData() == "project"


def write_variance_excel(rows, output_path, grouping_label, scope_label):
    """Writes the aggregated report as one styled sheet.

    Kept as a plain function (no Qt) so it can be exercised directly."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Variance"

    last_col = get_column_letter(len(COLUMNS))
    sheet.merge_cells(f"A1:{last_col}1")
    title = sheet.cell(row=1, column=1, value=f"{tr('var_title')} - {grouping_label}")
    title.font = Font(name="Calibri", bold=True, size=14)

    sheet.merge_cells(f"A2:{last_col}2")
    subtitle = sheet.cell(row=2, column=1, value=scope_label)
    subtitle.font = Font(name="Calibri", italic=True, size=10, color="666666")

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    header_font = Font(name="Calibri", bold=True, size=10)
    for index, label in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=4, column=index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = sheet.cell(row=5, column=1)

    for offset, data in enumerate(rows):
        excel_row = 5 + offset
        values = [
            data["group"], data["items"], data["delivered"],
            round(data["delivered_pct"] / 100.0, 4),
            data["awaiting_request"], data["attention"],
            _money(data["values"], data["currency"], data["mixed_currency"]),
        ]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=index, value=value)
            cell.font = Font(name="Calibri", size=10)
            if index == 4:
                cell.number_format = "0%"
            elif index > 1:
                cell.alignment = Alignment(horizontal="right")

    widths = (34, 10, 11, 13, 17, 17, 22)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(output_path)
    return output_path
