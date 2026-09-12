# -*- coding: utf-8 -*-
"""Excel export/import round-trip + PDF export (incl. Arabic content)."""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from database import Database
from export.excel_export import (
    export_project_to_excel, export_items_for_cost_update,
    export_flat_items_to_excel, export_projects_combined_to_excel,
)
from importers.excel_import import import_cost_updates_from_excel


def make_db():
    tmp = tempfile.mkdtemp(prefix="pt_xl_")
    db = Database(os.path.join(tmp, "db.sqlite"), os.path.join(tmp, "att"))
    return db, tmp


class TestExcelExport(unittest.TestCase):
    def setUp(self):
        self.db, self.tmp = make_db()
        self.pid = self.db.add_project("Excel Project", currency="SAR")
        self.a = self.db.add_item(self.pid, item_name="Pump A", total_quantity=4,
                                  unit_cost=2.5, currency="SAR")
        self.b = self.db.add_item(self.pid, item_name="Pump B", total_quantity=1,
                                  unit_cost=0, currency="SAR")

    def test_full_export_structure(self):
        path = os.path.join(self.tmp, "out.xlsx")
        export_project_to_excel(self.db, self.pid, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Excel Project")
        self.assertEqual(ws["A4"].value, "#")
        self.assertEqual(ws["B4"].value, "Item Name")
        # two data rows
        self.assertEqual(ws["B5"].value, "Pump A")
        self.assertEqual(ws["B6"].value, "Pump B")
        self.assertEqual(ws["G5"].value, 4)
        # TOTAL row has a real SUM formula (or 0 when empty)
        self.assertEqual(ws["L7"].value, "=SUM(L5:L6)")

    def test_filtered_export_only_includes_requested_ids(self):
        path = os.path.join(self.tmp, "filtered.xlsx")
        export_project_to_excel(self.db, self.pid, path, item_ids=[self.b])
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws["B5"].value, "Pump B")
        self.assertIsNone(ws["B6"].value)

    def test_empty_project_export_valid(self):
        pid = self.db.add_project("Empty")
        path = os.path.join(self.tmp, "empty.xlsx")
        export_project_to_excel(self.db, pid, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws["L5"].value, 0)  # no reversed SUM range

    def test_combined_flat_export(self):
        path = os.path.join(self.tmp, "combined.xlsx")
        export_projects_combined_to_excel(self.db, [self.pid], path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws["A2"].value, "Excel Project")
        self.assertEqual(ws["B3"].value, "Pump B")

    def test_flat_export_with_note(self):
        path = os.path.join(self.tmp, "flat.xlsx")
        export_flat_items_to_excel([], path, sheet_title="Search Results",
                                   header_note="Search: pump - 0 result(s)")
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.active["A1"].value, "Search: pump - 0 result(s)")


class TestExcelCostRoundTrip(unittest.TestCase):
    def test_round_trip_matched_by_item_id(self):
        db, tmp = make_db()
        pid = db.add_project("Costs")
        a = db.add_item(pid, item_name="A", unit_cost=0)
        b = db.add_item(pid, item_name="B", unit_cost=7)
        path = os.path.join(tmp, "costs.xlsx")
        export_items_for_cost_update(db, [pid], path)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # row 3 = item A: fill in a cost (was 0)
        ws.cell(row=3, column=8, value=12.5)
        # row 4 = item B: CLEAR the exported cost -> must be left untouched
        ws.cell(row=4, column=8).value = None
        # add a row for an item that no longer exists
        ws.cell(row=5, column=1, value=99999)
        ws.cell(row=5, column=8, value=3)
        wb.save(path)

        result = import_cost_updates_from_excel(db, path)
        self.assertEqual(result["updated"], 1)
        self.assertEqual(result["skipped_blank"], 1)
        self.assertEqual(result["not_found"], 1)
        self.assertAlmostEqual(db.get_item(a)["unit_cost"], 12.5)
        self.assertAlmostEqual(db.get_item(b)["unit_cost"], 7)  # untouched

    def test_import_rejects_foreign_file(self):
        db, tmp = make_db()
        path = os.path.join(tmp, "not_a_cost_file.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "hello"
        wb.save(path)
        with self.assertRaises(ValueError):
            import_cost_updates_from_excel(db, path)


class TestPdfExport(unittest.TestCase):
    def setUp(self):
        self.db, self.tmp = make_db()
        self.pid = self.db.add_project("\u0645\u0634\u0631\u0648\u0639 \u0627\u062e\u062a\u0628\u0627\u0631")  # Arabic project name
        self.db.add_item(self.pid, item_name="\u0645\u0636\u062e\u0629 \u063a\u0627\u0637\u0633\u0629 100",
                         total_quantity=2, unit_cost=5, currency="SAR")
        self.db.add_item(self.pid, item_name="Gate Valve", total_quantity=1, unit_cost=1)

    def test_project_pdf_builds_with_arabic(self):
        from export.pdf_export import export_project_to_pdf, _has_arabic, _shape_arabic
        self.assertTrue(_has_arabic("\u0645\u0636\u062e\u0629"))
        self.assertFalse(_has_arabic("Pump"))
        path = os.path.join(self.tmp, "report.pdf")
        export_project_to_pdf(self.db, self.pid, path)
        with open(path, "rb") as f:
            head = f.read(5)
        self.assertEqual(head, b"%PDF-")
        self.assertGreater(os.path.getsize(path), 4000)

    def test_dashboard_pdfs_build(self):
        from export.dashboard_pdf_export import (
            export_dashboard_summary_to_pdf, export_stale_items_to_pdf,
        )
        p1 = os.path.join(self.tmp, "dash.pdf")
        export_dashboard_summary_to_pdf(self.db, [self.pid], p1, "\u0645\u0634\u0631\u0648\u0639 \u0627\u062e\u062a\u0628\u0627\u0631")
        self.assertTrue(os.path.getsize(p1) > 3000)

        summary = self.db.get_dashboard_summary(stale_days=0)
        p2 = os.path.join(self.tmp, "stale.pdf")
        export_stale_items_to_pdf(self.db, summary["stale_items"], p2, "All Projects")
        self.assertTrue(os.path.getsize(p2) > 3000)




class TestPdfWithLogo(unittest.TestCase):
    def test_project_pdf_with_logo(self):
        import tempfile
        from unittest import mock
        import export.pdf_export as pdf_export
        tmp = tempfile.mkdtemp(prefix="pt_logo_pdf_")
        png = os.path.join(tmp, "logo.png")
        # a tiny valid PNG (1x1) via reportlab-free bytes
        import struct, zlib
        def chunk(typ, data):
            c = struct.pack(">I", len(data)) + typ + data
            return c + struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF)
        ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
        idat = zlib.compress(b"\x00\xff\x00\x00")
        with open(png, "wb") as f:
            f.write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b""))

        db = Database(os.path.join(tmp, "db.sqlite"), os.path.join(tmp, "att"))
        pid = db.add_project("Logo PDF")
        db.add_item(pid, item_name="Pump", total_quantity=1)
        out = os.path.join(tmp, "with_logo.pdf")
        with mock.patch.object(pdf_export._config, "logo_path", return_value=png):
            pdf_export.export_project_to_pdf(db, pid, out)
        self.assertTrue(os.path.getsize(out) > 4000)

    def test_project_pdf_without_logo(self):
        import tempfile
        from unittest import mock
        import export.pdf_export as pdf_export
        tmp = tempfile.mkdtemp(prefix="pt_nologo_pdf_")
        db = Database(os.path.join(tmp, "db.sqlite"), os.path.join(tmp, "att"))
        pid = db.add_project("No Logo PDF")
        db.add_item(pid, item_name="Pump", total_quantity=1)
        out = os.path.join(tmp, "no_logo.pdf")
        with mock.patch.object(pdf_export._config, "logo_path", return_value=None):
            pdf_export.export_project_to_pdf(db, pid, out)
        self.assertTrue(os.path.getsize(out) > 4000)


if __name__ == "__main__":
    unittest.main()
