"""
excel_export.py
Builds a professional-looking Excel workbook for one project, including
clickable hyperlinks to the first file of each attachment category
attached to each item (so a consultant or manager can click straight
through to the document) — both procurement documents (PR/PO/Reference)
and post-delivery/handover documents (Warranty, O&M Manual, Test
Certificate, As-Built Drawing, Spare Parts List, MIR, Invoice).
"""

import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from constants import (
    STATUS_COLORS, is_over_supplied, remaining_to_deliver, OVER_SUPPLY_COLOR,
    PRE_DELIVERY_CATEGORIES, POST_DELIVERY_CATEGORIES,
)

FONT_NAME = "Calibri"
HEADER_FILL = "1F4E78"
GROUP_FILL_PRE = "34506E"
GROUP_FILL_POST = "0E6B57"
TOTAL_FILL = "D9E1F2"

BASE_HEADERS = [
    "#", "Item Name", "Pump Station / Area", "Supplier", "Unit", "Currency",
    "Total Qty", "Requested Qty", "Delivered Qty", "Remaining to Deliver",
    "Unit Cost", "Total Cost", "Status", "Variance Note", "Remarks",
]
DOC_HEADERS = PRE_DELIVERY_CATEGORIES + POST_DELIVERY_CATEGORIES
N_BASE = len(BASE_HEADERS)
N_PRE = len(PRE_DELIVERY_CATEGORIES)
N_POST = len(POST_DELIVERY_CATEGORIES)


def _style_header(cell):
    cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_project_to_excel(db, project_id, output_path, item_ids=None):
    """Builds the full project Excel report. If item_ids is given, only
    those items are included (in that same relative order) — used for
    "Export Filtered View to Excel" so the file matches exactly what's
    currently visible after the search/status/supplier filters on the
    Items tab, no AI involved."""
    project = db.get_project(project_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Material Tracker"
    _write_project_sheet(ws, db, project, item_ids=item_ids)
    wb.save(output_path)
    return output_path


def export_multiple_projects_to_excel(db, project_ids, output_path):
    """Builds ONE workbook containing every given project, each on its
    own sheet/tab — used to export a whole (filtered) list of projects
    from the Projects screen in a single file instead of one file per
    project."""
    wb = Workbook()
    wb.remove(wb.active)  # start empty; we add exactly one sheet per project below
    used_titles = set()

    for project_id in project_ids:
        project = db.get_project(project_id)
        if project is None:
            continue
        title = _unique_sheet_title(project["name"], used_titles)
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        _write_project_sheet(ws, db, project)

    if not wb.sheetnames:
        wb.create_sheet(title="No Projects")

    wb.save(output_path)
    return output_path


_INVALID_SHEET_CHARS = set(':\\/?*[]')


def _unique_sheet_title(name, used_titles):
    """Excel sheet names must be <=31 chars and can't contain : \\ / ? * [ ] —
    and two projects can easily share a name/prefix once truncated to 31
    chars, so duplicates get a " (2)", " (3)", ... suffix instead of
    silently overwriting each other's tab."""
    cleaned = "".join(c for c in name if c not in _INVALID_SHEET_CHARS).strip() or "Project"
    base = cleaned[:31]
    title = base
    n = 2
    while title in used_titles:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    return title


FLAT_HEADERS = [
    "Project", "Item Name", "Pump Station / Area", "Supplier", "Unit", "Currency",
    "Total Qty", "Requested Qty", "Delivered Qty", "Unit Cost", "Total Cost",
    "Status", "Remarks",
]


COST_UPDATE_HEADERS = [
    "Item ID", "Project", "Item Name", "Pump Station / Area", "Supplier",
    "Unit", "Currency", "Unit Cost",
]


def export_items_for_cost_update(db, project_ids, output_path):
    """Exports every item from the given projects with an editable Unit
    Cost column, for a bulk price-filling workflow: fill in prices in
    Excel (much faster than opening each item individually), save, then
    use import_cost_updates_from_excel to apply them back. "Item ID" is
    the match key on re-import — don't edit or reorder that column.
    Leaving a Unit Cost cell blank on import means "don't change this
    item", so filling prices in over several sessions is safe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Update Costs"

    n_cols = len(COST_UPDATE_HEADERS)
    ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
    note = ws.cell(
        row=1, column=1,
        value="Fill in Unit Cost below, save, then use \u201cImport Costs from Excel\u201d. "
              "Leave a cell blank to skip that item (it won't be changed). Don't edit the Item ID column."
    )
    note.font = Font(name=FONT_NAME, italic=True, size=10, color="666666")

    header_row = 2
    for col, label in enumerate(COST_UPDATE_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        _style_header(cell)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    row_num = header_row + 1
    for project_id in project_ids:
        project = db.get_project(project_id)
        if project is None:
            continue
        for it in db.get_items(project_id):
            values = [
                it["id"], project["name"], it["item_name"], it["pump_station"] or "",
                it["supplier_name"] or "-", it["unit"] or "", it["currency"] or "-",
                it["unit_cost"] or 0,
            ]
            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=row_num, column=col, value=val)
                cell.font = Font(name=FONT_NAME, size=9 if col == 1 else 10,
                                  color="999999" if col == 1 else "000000")
                if col == n_cols:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
            row_num += 1

    widths = [10, 22, 34, 22, 18, 8, 9, 12]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{max(row_num - 1, header_row)}"

    wb.save(output_path)
    return output_path


def export_flat_items_to_excel(rows, output_path, sheet_title="Items", header_note=None):
    """Writes a single flat sheet — one row per item, with a Project
    column — instead of one-tab-per-project. Used for:
    - Search All Projects results ("export what I found")
    - several checked projects combined into one sheet for side-by-side
      comparison, instead of each on its own tab
    rows: an iterable of item-like dicts/Rows, each also carrying
    "project_name"."""
    wb = Workbook()
    ws = wb.active
    ws.title = _unique_sheet_title(sheet_title, set())

    start_row = 1
    if header_note:
        ws.merge_cells(f"A1:{get_column_letter(len(FLAT_HEADERS))}1")
        note_cell = ws.cell(row=1, column=1, value=header_note)
        note_cell.font = Font(name=FONT_NAME, italic=True, size=10, color="666666")
        start_row = 2

    header_row = start_row
    for col, label in enumerate(FLAT_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        _style_header(cell)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    row_num = header_row + 1
    for it in rows:
        total = it["total_quantity"] or 0
        unit_cost = it["unit_cost"] or 0
        values = [
            it["project_name"], it["item_name"], it["pump_station"] or "", it["supplier_name"] or "-",
            it["unit"] or "", it["currency"] or "-",
            total, it["requested_quantity"] or 0, it["delivered_quantity"] or 0,
            unit_cost, total * unit_cost, it["status"], it["remarks"] or "",
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            if col in (7, 8, 9, 10, 11):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
        row_num += 1

    widths = [22, 34, 22, 18, 8, 9, 11, 12, 12, 11, 12, 16, 30]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(FLAT_HEADERS))}{max(row_num - 1, header_row)}"

    wb.save(output_path)
    return output_path


STALE_HEADERS = ["Item Name", "Project", "Pump Station / Area", "Days Untouched"]


def export_stale_items_to_excel(stale_items, output_path, scope_label=None):
    """The "Possibly Forgotten" list, as its own standalone sheet —
    for sharing/saving separately from the rest of the Dashboard."""
    wb = Workbook()
    ws = wb.active
    ws.title = _unique_sheet_title("Possibly Forgotten", set())

    start_row = 1
    if scope_label:
        ws.merge_cells(f"A1:{get_column_letter(len(STALE_HEADERS))}1")
        note_cell = ws.cell(row=1, column=1, value=f"Scope: {scope_label}")
        note_cell.font = Font(name=FONT_NAME, italic=True, size=10, color="666666")
        start_row = 2

    header_row = start_row
    for col, label in enumerate(STALE_HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        _style_header(cell)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    row_num = header_row + 1
    for it in stale_items:
        values = [it["item_name"], it["project_name"], it["pump_station"] or "", it["days_stale"]]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.font = Font(name=FONT_NAME, size=10)
            if col == 4:
                cell.alignment = Alignment(horizontal="right")
        row_num += 1

    widths = [40, 24, 24, 15]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(STALE_HEADERS))}{max(row_num - 1, header_row)}"

    wb.save(output_path)
    return output_path


def export_projects_combined_to_excel(db, project_ids, output_path):
    """Every item from every given project, combined into ONE flat
    sheet (not one tab per project) — for checking/comparing several
    projects side by side in a single view."""
    rows = []
    for project_id in project_ids:
        project = db.get_project(project_id)
        if project is None:
            continue
        for it in db.get_items(project_id):
            row = dict(it)
            row["project_name"] = project["name"]
            rows.append(row)
    return export_flat_items_to_excel(rows, output_path, sheet_title="Combined")


def _write_project_sheet(ws, db, project, item_ids=None):
    """Writes one project's full report into an already-created
    worksheet — shared by both the single-project and multi-project
    (one-tab-per-project) export paths so their formatting never drifts
    apart."""
    project_id = project["id"]
    items = list(db.get_items(project_id))
    if item_ids is not None:
        wanted = set(item_ids)
        items = [it for it in items if it["id"] in wanted]

    all_attachments = db.get_all_attachments_for_project(project_id)
    att_by_item = defaultdict(list)
    for a in all_attachments:
        att_by_item[a["item_id"]].append(a)

    headers = BASE_HEADERS + DOC_HEADERS
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)

    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = project["name"]
    ws["A1"].font = Font(name=FONT_NAME, size=15, bold=True, color="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "Location:"
    ws["B2"] = project["location"] or ""
    ws["D2"] = "Contractor:"
    ws["E2"] = project["contractor"] or ""
    ws["G2"] = "Currency:"
    ws["H2"] = project["currency"] or ""
    for c in ("A2", "D2", "G2"):
        ws[c].font = Font(name=FONT_NAME, bold=True)

    # Group labels above the two attachment sections, so it's obvious at a
    # glance which columns are "before delivery" vs "after delivery" docs.
    group_row = 3
    pre_start_col = N_BASE + 1
    pre_end_col = N_BASE + N_PRE
    post_start_col = pre_end_col + 1
    post_end_col = post_start_col + N_POST - 1

    ws.merge_cells(start_row=group_row, start_column=pre_start_col, end_row=group_row, end_column=pre_end_col)
    pre_cell = ws.cell(row=group_row, column=pre_start_col, value="Procurement Documents")
    pre_cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    pre_cell.fill = PatternFill(start_color=GROUP_FILL_PRE, end_color=GROUP_FILL_PRE, fill_type="solid")
    pre_cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=group_row, start_column=post_start_col, end_row=group_row, end_column=post_end_col)
    post_cell = ws.cell(row=group_row, column=post_start_col, value="Post-Delivery / Handover Documents")
    post_cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    post_cell.fill = PatternFill(start_color=GROUP_FILL_POST, end_color=GROUP_FILL_POST, fill_type="solid")
    post_cell.alignment = Alignment(horizontal="center")

    header_row = 4
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=h)
        _style_header(cell)

    thin = Side(style="thin", color="C9CDD3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    data_font = Font(name=FONT_NAME, size=10)
    over_fill = PatternFill(start_color=OVER_SUPPLY_COLOR.lstrip("#"), end_color=OVER_SUPPLY_COLOR.lstrip("#"), fill_type="solid")

    row = header_row + 1
    for it in items:
        total = it["total_quantity"] or 0
        delivered = it["delivered_quantity"] or 0
        over = is_over_supplied(total, delivered)

        ws.cell(row=row, column=1, value=it["sort_order"])
        ws.cell(row=row, column=2, value=it["item_name"])
        ws.cell(row=row, column=3, value=it["pump_station"])
        ws.cell(row=row, column=4, value=it["supplier_name"] or "")
        ws.cell(row=row, column=5, value=it["unit"])
        ws.cell(row=row, column=6, value=it["currency"] or project["currency"])
        ws.cell(row=row, column=7, value=total)
        ws.cell(row=row, column=8, value=it["requested_quantity"])
        ws.cell(row=row, column=9, value=delivered)
        ws.cell(row=row, column=10, value=remaining_to_deliver(total, delivered))
        ws.cell(row=row, column=11, value=it["unit_cost"])
        ws.cell(row=row, column=12, value=f"=G{row}*K{row}")
        status_cell = ws.cell(row=row, column=13, value=("Over-supplied" if over else it["status"]))
        ws.cell(row=row, column=14, value=it["variance_note"] or "")
        ws.cell(row=row, column=15, value=it["remarks"])

        status_color = OVER_SUPPLY_COLOR.lstrip("#") if over else STATUS_COLORS.get(it["status"], "FFFFFF").lstrip("#")
        status_cell.fill = PatternFill(start_color=status_color, end_color=status_color, fill_type="solid")

        attachments = att_by_item.get(it["id"], [])
        for doc_idx, category in enumerate(DOC_HEADERS):
            col = N_BASE + 1 + doc_idx
            files = [a for a in attachments if a["category"] == category]
            _write_hyperlink_cell(ws, row, col, files)

        if over:
            for c in range(1, n_cols + 1):
                if c != 13:  # status already has its own fill
                    ws.cell(row=row, column=c).fill = over_fill

        row += 1

    last_data_row = row - 1
    for r in range(header_row + 1, last_data_row + 1):
        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font.name != FONT_NAME:
                cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for c in (7, 8, 9, 10, 11, 12):
            ws.cell(row=r, column=c).number_format = "#,##0.00"

    total_row = last_data_row + 1
    ws.cell(row=total_row, column=6, value="TOTAL").font = Font(name=FONT_NAME, bold=True)
    if last_data_row > header_row:
        total_value = f"=SUM(L{header_row + 1}:L{last_data_row})"
    else:
        # Empty project: a reversed range like SUM(L5:L4) is invalid —
        # write a plain 0 instead.
        total_value = 0
    total_cell = ws.cell(row=total_row, column=12, value=total_value)
    total_cell.number_format = "#,##0.00"
    for c in (6, 12):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = PatternFill(start_color=TOTAL_FILL, end_color=TOTAL_FILL, fill_type="solid")
        cell.border = border

    base_widths = {1: 5, 2: 40, 3: 20, 4: 16, 5: 8, 6: 9, 7: 10, 8: 11, 9: 11,
                   10: 13, 11: 10, 12: 12, 13: 15, 14: 22, 15: 22}
    for col, w in base_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for doc_idx, category in enumerate(DOC_HEADERS):
        col = N_BASE + 1 + doc_idx
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(16, len(category) + 2))

    ws.freeze_panes = "C5"
    ws.row_dimensions[header_row].height = 30
    ws.row_dimensions[group_row].height = 20


def _write_hyperlink_cell(ws, row, col, files):
    if not files:
        ws.cell(row=row, column=col, value="-")
        return
    first = files[0]
    cell = ws.cell(row=row, column=col)
    label = f"{len(files)} file(s)" if len(files) > 1 else "Open"
    cell.value = label
    cell.hyperlink = os.path.abspath(first["file_path"])
    cell.font = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")

