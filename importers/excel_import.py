"""
importers/excel_import.py
Bulk-imports items (and matches suppliers by name) from an existing Excel
spreadsheet into a project — lets someone switch from a spreadsheet-based
tracker to this app without retyping everything by hand.

Design: this module only does the mechanical part (read rows, parse them
per a given column mapping, write items). The actual field mapping is
always chosen/reviewed by the user in ui/import_excel_dialog.py — a
spreadsheet's column names never match ours exactly, and auto-guessing
silently would risk importing the wrong data into the wrong field.

Supplier names that don't already exist are left unmatched (item is
imported with no supplier set) rather than auto-created, so imports never
create a pile of messy near-duplicate supplier records from typos or
inconsistent spelling in someone's old spreadsheet.
"""

import openpyxl

IMPORT_FIELDS = [
    ("item_name", "Item Name", True),
    ("pump_station", "Pump Station / Area", False),
    ("supplier", "Supplier Name", False),
    ("unit", "Unit", False),
    ("currency", "Currency", False),
    ("total_quantity", "Total Quantity", False),
    ("requested_quantity", "Requested Quantity", False),
    ("delivered_quantity", "Delivered Quantity", False),
    ("unit_cost", "Unit Cost", False),
    ("remarks", "Remarks", False),
]

# Lower-cased keyword hints used only to *pre-select* a best-guess mapping
# in the UI — the user always sees and can override the guess before
# anything is imported.
_AUTO_MAP_HINTS = {
    "item_name": ["item name", "item", "description", "material", "name"],
    "pump_station": ["area", "location", "pump station", "station", "zone"],
    "supplier": ["supplier", "vendor", "manufacturer"],
    "unit": ["unit", "uom"],
    "currency": ["currency", "ccy"],
    "total_quantity": ["total qty", "total quantity", "boq qty", "qty", "quantity"],
    "requested_quantity": ["requested", "ordered qty", "po qty"],
    "delivered_quantity": ["delivered", "received qty"],
    "unit_cost": ["unit cost", "unit price", "rate", "price"],
    "remarks": ["remarks", "notes", "comment"],
}


def read_sheet_names(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def read_headers(file_path, sheet_name=None):
    """Returns (headers, header_row_index). Reads the first few rows and
    picks the one with the most non-empty text cells as the header row —
    tolerates a title row (e.g. a merged project-name cell) sitting above
    the real header row, which is common in hand-built BOQ sheets."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    best_row_idx, best_row, best_score = 1, [], -1
    max_scan = min(6, ws.max_row or 1)
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        score = sum(1 for c in row if isinstance(c, str) and c.strip())
        if score > best_score:
            best_score = score
            best_row_idx = row_idx
            best_row = row

    headers = [str(c).strip() if c is not None else "" for c in best_row]
    wb.close()
    return headers, best_row_idx


def guess_mapping(headers):
    """Best-effort auto mapping from detected headers to our field keys —
    purely a starting point the user reviews/edits in the UI, never
    applied silently."""
    mapping = {}
    used_headers = set()
    lower_headers = [(h, h.lower()) for h in headers if h]
    for field_key, _label, _required in IMPORT_FIELDS:
        hints = _AUTO_MAP_HINTS.get(field_key, [])
        for header, lower in lower_headers:
            if header in used_headers:
                continue
            if any(hint in lower for hint in hints):
                mapping[field_key] = header
                used_headers.add(header)
                break
    return mapping


def _parse_number(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return 0.0


def _row_to_record(row, col_index, mapping):
    record = {}
    for field_key, _label, _req in IMPORT_FIELDS:
        header = mapping.get(field_key)
        if header and header in col_index and col_index[header] < len(row):
            record[field_key] = row[col_index[header]]
        else:
            record[field_key] = None
    return record


def preview_rows(file_path, sheet_name, header_row_idx, mapping, limit=10):
    """Returns up to *limit* mapped rows (list of dicts) for a review
    table, without writing anything to the database."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    headers, _ = read_headers(file_path, sheet_name)
    col_index = {h: i for i, h in enumerate(headers)}

    rows = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if len(rows) >= limit:
            break
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        rows.append(_row_to_record(row, col_index, mapping))
    wb.close()
    return rows


def import_items_from_excel(db, project_id, file_path, sheet_name, header_row_idx, mapping):
    """Performs the actual import. Returns a result dict:
    {"imported": int, "skipped": [reasons...], "unmatched_suppliers": [names...]}."""
    project = db.get_project(project_id)
    suppliers_by_name = {s["name"].strip().lower(): s["id"] for s in db.get_suppliers()}

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    headers, _ = read_headers(file_path, sheet_name)
    col_index = {h: i for i, h in enumerate(headers)}

    imported = 0
    skipped = []
    unmatched_suppliers = set()

    for row_num, row in enumerate(
        ws.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1
    ):
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        record = _row_to_record(row, col_index, mapping)

        item_name = record["item_name"]
        item_name = str(item_name).strip() if item_name is not None else ""
        if not item_name:
            skipped.append(f"Row {row_num}: no item name")
            continue

        supplier_id = None
        supplier_raw = record["supplier"]
        if supplier_raw:
            key = str(supplier_raw).strip().lower()
            if key in suppliers_by_name:
                supplier_id = suppliers_by_name[key]
            else:
                unmatched_suppliers.add(str(supplier_raw).strip())

        currency = record["currency"]
        currency = str(currency).strip() if currency else project["currency"]

        db.add_item(
            project_id,
            item_name=item_name,
            pump_station=str(record["pump_station"] or "").strip(),
            supplier_id=supplier_id,
            unit=str(record["unit"] or "").strip(),
            currency=currency,
            total_quantity=_parse_number(record["total_quantity"]),
            requested_quantity=_parse_number(record["requested_quantity"]),
            delivered_quantity=_parse_number(record["delivered_quantity"]),
            unit_cost=_parse_number(record["unit_cost"]),
            remarks=str(record["remarks"] or "").strip(),
        )
        imported += 1

    wb.close()
    return {
        "imported": imported,
        "skipped": skipped,
        "unmatched_suppliers": sorted(unmatched_suppliers),
    }


def import_cost_updates_from_excel(db, file_path):
    """Reads back a file produced by export_items_for_cost_update and
    applies Unit Cost changes by matching on the "Item ID" column —
    never by name, so it's exact even across projects with similarly
    named items. A blank Unit Cost cell means "leave this item alone",
    which makes filling prices in over several sessions safe (you never
    have to finish the whole sheet in one sitting). Returns a summary
    dict: {"updated", "skipped_blank", "not_found"}."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    header_row_idx = None
    header_values = None
    max_scan = min(6, ws.max_row or 1)
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        if row and str(row[0] or "").strip().lower() == "item id":
            header_row_idx = idx
            header_values = row
            break
    if header_row_idx is None:
        wb.close()
        raise ValueError(
            "Could not find the \u201cItem ID\u201d header row \u2014 is this a cost-update "
            "file exported by this app's \u201cExport Items for Cost Update\u201d?"
        )

    id_col = 0  # "Item ID" is always column A in files this app produces
    try:
        cost_col = list(header_values).index("Unit Cost")
    except ValueError:
        wb.close()
        raise ValueError("Could not find the \u201cUnit Cost\u201d column in this file.")

    updated = 0
    skipped_blank = 0
    not_found = 0

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        raw_id = row[id_col] if id_col < len(row) else None
        if raw_id is None or str(raw_id).strip() == "":
            continue
        try:
            item_id = int(raw_id)
        except (TypeError, ValueError):
            continue

        raw_cost = row[cost_col] if cost_col < len(row) else None
        if raw_cost is None or str(raw_cost).strip() == "":
            skipped_blank += 1
            continue
        try:
            cost_value = float(str(raw_cost).replace(",", "").strip())
        except (TypeError, ValueError):
            skipped_blank += 1
            continue

        existing = db.get_item(item_id)
        if existing is None or existing["is_deleted"]:
            not_found += 1
            continue

        db.update_item(item_id, unit_cost=cost_value)
        updated += 1

    wb.close()
    return {"updated": updated, "skipped_blank": skipped_blank, "not_found": not_found}

